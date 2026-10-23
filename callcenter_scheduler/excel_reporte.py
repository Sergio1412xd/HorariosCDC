"""
excel_reporte.py
Exporta el horario generado a un archivo Excel formateado -- útil para integrarlo con
nómina o para que un supervisor lo edite/imprima directamente, algo estándar en cualquier
sistema profesional de Workforce Management.
"""

from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db import DIAS_SEMANA
from generador_horarios import generar_horario_semana, estimar_nivel_servicio_real
from costos import calcular_costo_horario
from cumplimiento import validar_cumplimiento
from db import listar_empleados

INK = "1B2430"
ROJO_KFC = "C8102E"
GRIS_CLARO = "F3F3F1"

FUENTE_TITULO = Font(name="Arial", bold=True, size=15, color=INK)
FUENTE_HEADER = Font(name="Arial", bold=True, color="FFFFFF", size=10)
FUENTE_NORMAL = Font(name="Arial", size=10)
FUENTE_NEGRITA = Font(name="Arial", bold=True, size=10)
RELLENO_HEADER = PatternFill("solid", fgColor=INK)
RELLENO_ALTERNO = PatternFill("solid", fgColor=GRIS_CLARO)
BORDE_FINO = Border(*[Side(style="thin", color="E1E3DF")] * 4)

MOTIVOS_TEXTO = {
    "no_disponible": "No disponible", "sin_informacion": "Sin información",
    "libre_no_requerido": "Libre", "tope_semanal_alcanzado": "Tope semanal",
    "descanso_asignado": "Descanso",
}


def generar_excel_horario(fecha_lunes: date, target_service_level: float = 0.95, shrinkage: float = 0.10,
                           max_occupancy: float = 0.85, semanas_base: list = None, mostrar_costos: bool = True,
                           db_path: str = "callcenter.db", salida: str = "horario.xlsx",
                           resultado: dict = None, costo: dict = None, nivel_servicio: dict = None) -> str:
    """Si se pasan 'resultado'/'costo'/'nivel_servicio' ya calculados (horario tal cual está
    en pantalla -- guardado del historial o con ediciones manuales), el Excel se arma con
    esos datos sin volver a generar el horario desde cero (ver el mismo parámetro en
    pdf_reporte.generar_pdf_horario para el porqué)."""
    empleados = listar_empleados(db_path=db_path)
    if resultado is None:
        resultado = generar_horario_semana(fecha_lunes, db_path=db_path, target_service_level=target_service_level,
                                            shrinkage=shrinkage, max_occupancy=max_occupancy, semanas_base=semanas_base)
    if costo is None:
        costo = calcular_costo_horario(resultado, db_path=db_path)
    if nivel_servicio is None:
        nivel_servicio = estimar_nivel_servicio_real(resultado)
    cumplimiento = validar_cumplimiento(resultado, fecha_lunes, empleados, db_path=db_path)
    horas_trabajadas = [h for h in resultado["horas_semana"].values() if h > 0]
    promedio_horas = round(sum(horas_trabajadas) / len(horas_trabajadas), 1) if horas_trabajadas else 0

    wb = openpyxl.Workbook()

    # ================= Hoja 1: Cuadrante de turnos =================
    ws = wb.active
    ws.title = "Horario"

    ws["A1"] = "Cuadrante de Turnos — Call Center"
    ws["A1"].font = FUENTE_TITULO
    ws["A2"] = f"Semana del {fecha_lunes.isoformat()}  |  Meta de servicio: {int(target_service_level*100)}%"
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="6B7280")

    fila_header = 4
    headers = ["Asesor"] + DIAS_SEMANA + ["Horas/semana"] + (["Costo semana"] if mostrar_costos else [])
    n_cols = len(headers)
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=fila_header, column=col, value=h)
        c.font = FUENTE_HEADER
        c.fill = RELLENO_HEADER
        c.alignment = Alignment(horizontal="center")
        c.border = BORDE_FINO

    fila = fila_header + 1
    for e in empleados:
        nombre = e["nombre"]
        horario_emp = resultado["horario"].get(nombre, {})
        relleno = RELLENO_ALTERNO if fila % 2 == 0 else None

        ws.cell(row=fila, column=1, value=nombre).font = FUENTE_NEGRITA
        for j, dia in enumerate(DIAS_SEMANA, start=2):
            info = horario_emp.get(dia, {"motivo": "?"})
            if "turnos" in info:
                def _texto_turno_xlsx(t):
                    base = f"{t['inicio']:02d}:00-{t['fin']:02d}:00{' (exc.)' if t.get('excepcion') else ''}"
                    if t.get("almuerzo_horas"):
                        ai, af = t["almuerzo_inicio"], t["almuerzo_inicio"] + t["almuerzo_horas"]
                        base += f" [almuerzo {ai:02d}:00-{af:02d}:00]"
                    return base
                texto = " + ".join(_texto_turno_xlsx(t) for t in info["turnos"])
            else:
                texto = MOTIVOS_TEXTO.get(info.get("motivo"), info.get("motivo", ""))
                if str(info.get("motivo", "")).startswith("permiso:"):
                    texto = info["motivo"].replace("permiso:", "").capitalize()
            c = ws.cell(row=fila, column=j, value=texto)
            c.font = FUENTE_NORMAL
            c.alignment = Alignment(horizontal="center")

        ws.cell(row=fila, column=9, value=resultado["horas_semana"].get(nombre, 0)).font = FUENTE_NORMAL
        if mostrar_costos:
            ws.cell(row=fila, column=10, value=round(costo["costo_por_empleado"].get(nombre, 0))).font = FUENTE_NORMAL
            ws.cell(row=fila, column=10).number_format = '$#,##0'

        if relleno:
            for col in range(1, n_cols + 1):
                if ws.cell(row=fila, column=col).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(row=fila, column=col).fill = relleno
        for col in range(1, n_cols + 1):
            ws.cell(row=fila, column=col).border = BORDE_FINO
        fila += 1

    # Fila de totales
    ws.cell(row=fila, column=1, value="TOTAL").font = Font(name="Arial", bold=True, color=ROJO_KFC)
    ws.cell(row=fila, column=9, value=round(sum(resultado["horas_semana"].values()))).font = Font(name="Arial", bold=True)
    if mostrar_costos:
        ws.cell(row=fila, column=10, value=round(costo["total_semana"])).font = Font(name="Arial", bold=True, color=ROJO_KFC)
        ws.cell(row=fila, column=10).number_format = '$#,##0'

    ws.column_dimensions["A"].width = 24
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["I"].width = 13
    if mostrar_costos:
        ws.column_dimensions["J"].width = 14
    ws.freeze_panes = "B5"

    # ================= Hoja 2: Resumen =================
    ws2 = wb.create_sheet("Resumen")
    ws2["A1"] = "Resumen Ejecutivo"
    ws2["A1"].font = FUENTE_TITULO

    filas_resumen = [
        ("Semana", fecha_lunes.isoformat()),
        ("Meta de servicio", f"{int(target_service_level*100)}%"),
        ("Shrinkage", f"{int(shrinkage*100)}%"),
        ("% Contestación estimado", f"{nivel_servicio['sl_estimado_semana_pct']}%"),
        ("Horas-persona programadas", round(sum(resultado["horas_semana"].values()))),
        ("Horas/semana promedio por asesor", promedio_horas),
    ]
    if mostrar_costos:
        filas_resumen.insert(4, ("Costo total semana", round(costo["total_semana"])))
    filas_resumen += [
        ("Alertas de cobertura", len([a for a in resultado["alertas"] if "faltan" in a])),
        ("Cumplimiento normativo", "OK" if cumplimiento["resumen"]["cumple"] else f"{cumplimiento['resumen']['total_violaciones']} violación(es)"),
        ("Máx. días consecutivos detectado", cumplimiento["resumen"]["dias_consecutivos_max_detectado"]),
    ]
    for i, (label, valor) in enumerate(filas_resumen, start=3):
        ws2.cell(row=i, column=1, value=label).font = FUENTE_NEGRITA
        c = ws2.cell(row=i, column=2, value=valor)
        c.font = FUENTE_NORMAL
        if label == "Costo total semana":
            c.number_format = '$#,##0'
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 20

    if cumplimiento["violaciones"]:
        ws2.cell(row=len(filas_resumen) + 5, column=1, value="Violaciones detectadas:").font = Font(name="Arial", bold=True, color=ROJO_KFC)
        for i, v in enumerate(cumplimiento["violaciones"], start=len(filas_resumen) + 6):
            ws2.cell(row=i, column=1, value=v["detalle"]).font = Font(name="Arial", size=9)

    wb.save(salida)
    return salida


if __name__ == "__main__":
    import sys
    from datetime import timedelta
    lunes = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today()
    ruta = generar_excel_horario(lunes)
    print(f"Excel generado: {ruta}")
